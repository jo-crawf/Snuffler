#!/usr/bin/env python3
"""Conformance suite for Snuffler.

GhostBuster's suite ran its two Windows builds side by side and demanded
byte-identical media, because both drove the same WIC encoders. There is no WIC
on a Mac, so Snuffler is held to what that byte-identity stood for:

  * every structural check GhostBuster's suite makes, unchanged
  * the same decisions wherever a decision is logic rather than encoder:
    what was rejected, ghosts removed, pictures kept, media removed and
    layout-data sizes, all field for field
  * media left alone are byte-identical to GhostBuster's
  * every image GhostBuster shrinks, Snuffler shrinks too, under the same new
    name and at the same pixel dimensions. It may shrink a few more, where
    its encoder beats the original and WIC's did not
  * the lossless tier is pixel-identical, both to GhostBuster and to the
    source; Med stays within a small mean difference of GhostBuster; XSmall
    ends up no further from the source than GhostBuster's JPEG, give or take
    one level in 255

The reference is GhostBuster's Rust build. On Windows it can run live
(--reference path/to/ghostbuster.exe); everywhere else the suite compares
against conformance/golden/, which was produced from it. Regenerate that with
--write-golden whenever the fixtures change.

    python conformance/run.py                               # Snuffler vs golden
    python conformance/run.py --bin path/to/Snuffler        # a specific binary
    python conformance/run.py --arch x86_64                 # Intel slice, via Rosetta
    python conformance/run.py --reference ghostbuster.exe   # live, Windows only
    python conformance/run.py --reference ghostbuster.exe --write-golden
    python conformance/run.py --fixtures DIR --reference ghostbuster.exe
                                  # private workbooks; never commit those

Exit code is 0 only when every case passes.
"""

import argparse
import io
import json
import os
import posixpath
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIXTURES = os.path.join(HERE, "fixtures")
GOLDEN = os.path.join(HERE, "golden")

# Each case is (name, bust?, squish quality or None).
CASES = [
    ("bust", True, None),
    ("squish-high", False, "High"),
    ("squish-med", False, "Med"),
    ("squish-xsmall", False, "XSmall"),
    ("bust+xsmall", True, "XSmall"),
]

# Fields that are pure logic and must match GhostBuster exactly. Whether an
# image was converted, and so whether a squish-only run changed anything,
# depends on the encoder too; those are checked image by image instead.
EXACT_FIELDS = [
    "rejected",
    "ghosts_removed",
    "pictures_kept",
    "media_removed",
    "drawing_before",
    "drawing_after",
    "file_before",
]

# Med: largest mean absolute difference per channel (0-255) between the two
# downsampled images. WIC's Fant filter and the triangle filter used here never
# agree exactly; within this they are the same picture.
MED_LIMIT = 3.0
# XSmall: how much further from the source Snuffler's JPEG may be than
# GhostBuster's, in the same units.
XSMALL_SLACK = 1.0
# A JPEG decoded by two different decoders can differ by a level here and
# there, so "lossless" from a JPEG source allows that much.
JPEG_DECODE_SLACK = 1.0


def log(msg=""):
    print(msg, flush=True)


def default_bin():
    if os.environ.get("SNUFFLER_BIN"):
        return os.environ["SNUFFLER_BIN"]
    for cand in (
        os.path.join(ROOT, "target", "release", "snuffler.exe"),
        os.path.join(ROOT, "target", "release", "snuffler"),
        os.path.join(ROOT, "dist", "Snuffler.app", "Contents", "MacOS", "Snuffler"),
    ):
        if os.path.exists(cand):
            return cand
    return None


def run_tool(prefix, exe, files, bust, quality, workdir, tag):
    out = os.path.join(workdir, tag + ".json")
    args = prefix + [exe, "--headless", "--json", out]
    if bust:
        args.append("--bust")
    if quality:
        args += ["--squish", quality]
    args += files
    # GhostBuster is a windows-subsystem binary that detaches from the
    # console; run() waits on the process itself, so both are waited for.
    proc = subprocess.run(args, capture_output=True, text=True)
    if proc.stderr.strip():
        log(proc.stderr.strip())
    return _load(out)


def _load(path):
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8-sig") as fh:
        return json.load(fh)


def output_name(source):
    stem, ext = os.path.splitext(source)
    return "%s (cleaned)%s" % (stem, ext)


# ------------------------------------------------------------ structural
# Identical to GhostBuster's conformance/run.py.

def structural_checks(xlsx, original, busted):
    """Checks that hold for any output, regardless of implementation.

    `busted` says whether this case asked for ghost removal. Squish-only runs
    are expected to leave ghosts in place -- asserting otherwise would be
    testing for something the caller never requested.
    """
    problems = []
    try:
        z = zipfile.ZipFile(xlsx)
    except Exception as exc:
        return ["not a readable zip: %s" % exc]
    names = set(z.namelist())
    if z.testzip() is not None:
        problems.append("zip integrity failed")

    # Every internal relationship target resolves to a real part.
    for n in sorted(names):
        if not n.endswith(".rels"):
            continue
        part_dir = posixpath.dirname(posixpath.dirname(n))
        body = z.read(n).decode("utf-8", "replace")
        for tag in re.findall(r"<Relationship\b[^>]*/>", body):
            if 'TargetMode="External"' in tag:
                continue
            m = re.search(r'Target="([^"]+)"', tag)
            if not m:
                continue
            target = posixpath.normpath(posixpath.join(part_dir, m.group(1)))
            if target not in names:
                problems.append("dangling relationship %s -> %s" % (n, m.group(1)))

    # Every media extension is declared in [Content_Types].
    ct = z.read("[Content_Types].xml").decode("utf-8", "replace")
    for ext in {n.rsplit(".", 1)[-1].lower() for n in names if n.startswith("xl/media/")}:
        if 'Extension="%s"' % ext not in ct:
            problems.append("media type %s not declared" % ext)

    # richData (in-cell images) must survive untouched.
    zo = zipfile.ZipFile(original)
    want = len([n for n in zo.namelist() if "richData" in n])
    got = len([n for n in names if "richData" in n])
    if want != got:
        problems.append("richData parts %d -> %d" % (want, got))

    # After a bust, no zero-area anchor may remain in any drawing.
    if busted:
        for n in names:
            if not re.match(r"xl/drawings/drawing\d+\.xml$", n):
                continue
            xml = z.read(n).decode("utf-8", "replace")
            for m in re.finditer(
                r"(?s)<xdr:(oneCellAnchor|absoluteAnchor)\b[^>]*>.*?</xdr:\1>", xml
            ):
                ext = re.search(r'<xdr:ext\s+cx="(\d+)"\s+cy="(\d+)"', m.group(0))
                if ext and (ext.group(1) == "0" or ext.group(2) == "0"):
                    problems.append("ghost anchor survived in %s" % n)
                    break
    return problems


def cells_match(a, b):
    try:
        import openpyxl
    except ImportError:
        return None
    import warnings
    with warnings.catch_warnings():
        # openpyxl announces every WMF/EMF it cannot render; only cell
        # values matter here.
        warnings.simplefilter("ignore")
        wa = openpyxl.load_workbook(a, data_only=False)
        wb = openpyxl.load_workbook(b, data_only=False)
    if wa.sheetnames != wb.sheetnames:
        return False
    for s in wb.sheetnames:
        for row in wb[s].iter_rows():
            for c in row:
                if wa[s][c.coordinate].value != c.value:
                    return False
    return True


# ---------------------------------------------------------------- pixels

def emf_bitmap(data):
    """The bitmap inside a HEADER + STRETCHDIBITS + EOF metafile, or None.
    Pillow can only render EMF on Windows, so this reads it directly."""
    from PIL import Image
    recs, off = [], 0
    while off + 8 <= len(data):
        t, sz = struct.unpack_from("<II", data, off)
        if sz < 8:
            return None
        recs.append((t, off))
        off += sz
    if [r[0] for r in recs] != [1, 81, 14]:
        return None
    ro = recs[1][1]
    off_bmi, _, off_bits, _ = struct.unpack_from("<IIII", data, ro + 48)
    w, h = struct.unpack_from("<ii", data, ro + off_bmi + 4)
    raw = data[ro + off_bits: ro + off_bits + w * abs(h) * 4]
    return Image.frombytes("RGB", (w, abs(h)), raw, "raw", "BGRX", 0, -1 if h > 0 else 1)


def decode(name, data):
    from PIL import Image
    if name.lower().endswith((".emf", ".wmf")):
        return emf_bitmap(data)
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
        return im
    except Exception:
        return None


def mad(a, b):
    from PIL import ImageChops, ImageStat
    diff = ImageChops.difference(a.convert("RGB"), b.convert("RGB"))
    return sum(ImageStat.Stat(diff).mean) / 3.0


def same_pixels(a, b, from_jpeg):
    if a.size != b.size:
        return False
    if a.convert("RGBA").tobytes() == b.convert("RGBA").tobytes():
        return True
    return from_jpeg and mad(a, b) <= JPEG_DECODE_SLACK


# ------------------------------------------------------------ comparison

def zip_media(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist() if n.startswith("xl/media/")}


def stem(name):
    return name.rsplit(".", 1)[0]


class Tally:
    def __init__(self):
        self.identical = self.exact = self.close = self.extra = 0
        self.worst = 0.0
        self.ref_bytes = self.our_bytes = 0


def compare(ref, ours, ref_dir, our_dir, snap, quality):
    """Field-by-field agreement with GhostBuster, then image by image."""
    problems = []
    t = Tally()
    if len(ref["results"]) != len(ours["results"]):
        return ["result count %d vs %d" % (len(ref["results"]), len(ours["results"]))], t

    for r, o in zip(ref["results"], ours["results"]):
        src = r.get("source")
        if src != o.get("source"):
            problems.append("source order %s vs %s" % (src, o.get("source")))
            continue
        fields = EXACT_FIELDS + (["changed", "images_converted"] if quality is None else [])
        for f in fields:
            if r.get(f) != o.get(f):
                problems.append("%s: %s %r != %r" % (src, f, r.get(f), o.get(f)))
        if r.get("rejected") or r.get("error"):
            continue

        source = zip_media(os.path.join(snap, src))
        theirs = zip_media(os.path.join(ref_dir, output_name(src))) if r.get("changed") else source
        mine = zip_media(os.path.join(our_dir, output_name(src))) if o.get("changed") else source
        by_src = {stem(n): n for n in source}
        by_ref = {stem(n): n for n in theirs}
        by_our = {stem(n): n for n in mine}
        if set(by_ref) != set(by_our):
            problems.append("%s: different media survived (GhostBuster only %s / Snuffler only %s)"
                            % (src, sorted(set(by_ref) - set(by_our)), sorted(set(by_our) - set(by_ref))))
            continue

        for k in sorted(by_ref):
            sn, rn, on = by_src[k], by_ref[k], by_our[k]
            s, rb, ob = source[sn], theirs[rn], mine[on]
            ref_conv = rn != sn or rb != s
            our_conv = on != sn or ob != s
            label = "%s: %s" % (src, sn)
            from_jpeg = sn.lower().endswith((".jpg", ".jpeg"))

            if not ref_conv and not our_conv:
                t.identical += 1
                continue
            if quality is None:
                problems.append("%s re-encoded, but this case never squishes" % label)
                continue
            if ref_conv and not our_conv:
                problems.append("%s: GhostBuster shrank it (%d -> %d bytes), Snuffler left it"
                                % (label, len(s), len(rb)))
                continue

            b = decode(on, ob)
            if b is None:
                problems.append("%s: Snuffler's output does not decode" % label)
                continue

            if not ref_conv:
                # Snuffler's encoder beat the original where WIC's did not.
                # Allowed, but the lossless tier must still be lossless.
                t.extra += 1
                if quality == "High":
                    orig = decode(sn, s)
                    if orig is not None and not same_pixels(orig, b, from_jpeg):
                        problems.append("%s: extra lossless conversion changed pixels" % label)
                continue

            if on != rn:
                problems.append("%s: renamed to %s, GhostBuster used %s" % (label, on, rn))
            a = decode(rn, rb)
            if a is None:
                problems.append("%s: GhostBuster's output does not decode" % label)
                continue
            t.ref_bytes += len(rb)
            t.our_bytes += len(ob)
            if a.size != b.size:
                problems.append("%s: %dx%d, GhostBuster made %dx%d"
                                % (label, b.size[0], b.size[1], a.size[0], a.size[1]))
                continue
            if quality == "High":
                if same_pixels(a, b, from_jpeg):
                    t.exact += 1
                else:
                    problems.append("%s: lossless, but pixels differ (MAD %.3f)" % (label, mad(a, b)))
            elif quality == "Med":
                m = mad(a, b)
                t.worst = max(t.worst, m)
                if m <= MED_LIMIT:
                    t.close += 1
                else:
                    problems.append("%s: differs from GhostBuster by MAD %.2f (limit %.1f)"
                                    % (label, m, MED_LIMIT))
            else:
                orig = decode(sn, s)
                if orig is None or orig.size != b.size:
                    m, limit = mad(a, b), MED_LIMIT
                    what = "from GhostBuster"
                else:
                    theirs_off = mad(a, orig)
                    m, limit = mad(b, orig), theirs_off + XSMALL_SLACK
                    what = "from the source (GhostBuster's is %.2f)" % theirs_off
                if m <= limit:
                    t.close += 1
                else:
                    problems.append("%s: MAD %.2f %s" % (label, m, what))
    return problems, t


# ------------------------------------------------------------------ main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bin", default=default_bin(), help="Snuffler binary (headless-capable)")
    ap.add_argument("--arch", help="run the binary under `arch -<arch>` (macOS)")
    ap.add_argument("--reference", help="GhostBuster's ghostbuster.exe, to compare live")
    ap.add_argument("--write-golden", action="store_true",
                    help="record --reference's results as conformance/golden/")
    ap.add_argument("--fixtures", default=FIXTURES)
    ap.add_argument("--keep", action="store_true", help="leave the work directory")
    args = ap.parse_args()

    if not os.path.isdir(args.fixtures):
        log("no fixtures at %s" % args.fixtures)
        return 1
    # Never treat our own output as input: a "(cleaned)" file left in the
    # fixtures folder would quietly become a test case for the next run.
    fixtures = sorted(
        f for f in os.listdir(args.fixtures)
        if f.endswith((".xlsx", ".xlsm")) and "(cleaned" not in f
    )
    if not fixtures:
        log("no fixtures found")
        return 1

    same_fixtures = os.path.abspath(args.fixtures) == os.path.abspath(FIXTURES)
    if args.write_golden:
        if not args.reference:
            log("--write-golden needs --reference")
            return 1
        if not same_fixtures:
            log("golden results are only ever recorded from the committed fixtures")
            return 1
    elif not args.bin or not os.path.exists(args.bin):
        log("Snuffler binary not found: %s (build it, or pass --bin)" % args.bin)
        return 1
    elif not args.reference and not os.path.isdir(GOLDEN):
        log("no golden results at %s and no --reference to run" % GOLDEN)
        return 1
    elif not args.reference and not same_fixtures:
        log("golden results only cover the committed fixtures; pass --reference for others")
        return 1

    prefix = ["arch", "-" + args.arch] if args.arch else []
    log("fixtures: %d      cases: %d      reference: %s%s"
        % (len(fixtures), len(CASES), "live" if args.reference else "golden",
           "      arch: " + args.arch if args.arch else ""))
    log("=" * 78)

    root = tempfile.mkdtemp(prefix="snufconf_")
    snap = os.path.join(root, "_fixtures")
    os.makedirs(snap)
    for f in fixtures:
        shutil.copy2(os.path.join(args.fixtures, f), os.path.join(snap, f))

    def fresh_copies(sub):
        os.makedirs(sub)
        copies = []
        for f in fixtures:
            dst = os.path.join(sub, f)
            shutil.copy2(os.path.join(snap, f), dst)
            copies.append(dst)
        return copies

    if args.write_golden:
        shutil.rmtree(GOLDEN, ignore_errors=True)
        for case, bust, quality in CASES:
            sub = os.path.join(root, case)
            data = run_tool([], args.reference, fresh_copies(sub), bust, quality, sub, "reference")
            if data is None:
                log("%-14s reference emitted no JSON" % case)
                return 1
            dest = os.path.join(GOLDEN, case)
            os.makedirs(dest)
            for rec in data["results"]:
                if rec.get("changed"):
                    shutil.copy2(os.path.join(sub, output_name(rec["source"])), dest)
            with open(os.path.join(GOLDEN, case + ".json"), "w", encoding="utf-8") as fh:
                json.dump(data, fh, indent=1, sort_keys=True)
            log("%-14s recorded %d results" % (case, len(data["results"])))
        shutil.rmtree(root, ignore_errors=True)
        return 0

    failures = 0
    checked = 0
    for case, bust, quality in CASES:
        work = os.path.join(root, case)
        ours_dir = os.path.join(work, "snuffler")
        ours = run_tool(prefix, args.bin, fresh_copies(ours_dir), bust, quality, ours_dir, "snuffler")
        if ours is None:
            log("%-14s NO JSON EMITTED" % case)
            failures += 1
            continue

        # Structural checks apply to whatever Snuffler produced.
        for rec in ours["results"]:
            if not rec.get("changed"):
                continue
            out = os.path.join(ours_dir, output_name(rec["source"]))
            if not os.path.exists(out):
                log("%-14s missing output for %s" % (case, rec["source"]))
                failures += 1
                continue
            checked += 1
            probs = structural_checks(out, os.path.join(snap, rec["source"]), bust)
            if cells_match(out, os.path.join(snap, rec["source"])) is False:
                probs.append("cell values changed")
            for p in probs:
                log("%-14s %s: %s" % (case, rec["source"], p))
                failures += 1

        if args.reference:
            ref_dir = os.path.join(work, "reference")
            ref = run_tool([], args.reference, fresh_copies(ref_dir), bust, quality, ref_dir, "reference")
        else:
            ref_dir = os.path.join(GOLDEN, case)
            ref = _load(os.path.join(GOLDEN, case + ".json"))
        if ref is None:
            log("%-14s no reference results" % case)
            failures += 1
            continue

        probs, t = compare(ref, ours, ref_dir, ours_dir, snap, quality)
        parts = ["untouched %d" % t.identical]
        if t.exact:
            parts.append("pixel-exact %d" % t.exact)
        if t.close:
            parts.append("close %d" % t.close)
        if t.extra:
            parts.append("+%d more shrunk" % t.extra)
        if t.ref_bytes:
            parts.append("bytes %+.1f%%" % (100.0 * (t.our_bytes - t.ref_bytes) / t.ref_bytes))
        log("%-14s %-8s  %s" % (case, "OK" if not probs else "MISMATCH", ", ".join(parts)))
        for p in probs:
            log("    %s" % p)
            failures += 1

    if args.keep:
        log("\nwork kept at %s" % root)
    else:
        shutil.rmtree(root, ignore_errors=True)

    log("=" * 78)
    log("%d outputs structurally checked; %d problems" % (checked, failures))
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
